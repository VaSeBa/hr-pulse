"""Запись выдуманной компании в Excel в формате выгрузок 1С:ЗУП.

Файлы повторяют особенности настоящих выгрузок: шапку отчёта над таблицей,
пустые «служебные» колонки, двухуровневые заголовки, даты текстом.
На них мы научимся читать реальные выгрузки, не трогая реальные ПДн.
"""

from pathlib import Path

from openpyxl import Workbook

from hr_pulse.core.formats_1c import format_count, ru_date
from hr_pulse.demo.generator import REASONS, TECHNICAL_REASON, FakeRow

ORGANIZATION = "Демо-Строй"

# Шапка таблицы «Личные данные сотрудников». None — пустая «служебная» колонка,
# как в настоящей выгрузке 1С, где часть колонок объединена и стоит пустой.
EMPLOYEES_HEADER = (
    "Табельный номер (с префиксами)",
    None,
    None,
    "ФИО",
    None,
    "Должность",
    "Дата приема",
    "Дата увольнения",  # в настоящей выгрузке у этой колонки был неверный заголовок
    "Должность по штатному расписанию.Подкатегория персонала",
    "СНИЛС",
    "Подразделение",
    "Подразделение.Подразделение верхнего уровня (для отчетов) (Подразделения)",
    "Подразделение.Функциональное подчинение (Подразделения)",
    "Подразделение.Функциональный блок (Подразделения)",
)
EMPLOYEES_HEADER_ROW = 9  # строка заголовков, как в настоящей выгрузке
MISLABELED_FIRE_DATE = "Адрес места проживания"

# Причины, которые НЕ входят в учётную текучесть (как в справочнике компании)
EXCLUDED_REASONS = {
    "Несоблюдение ПВТР",
    "Истечение срока патента, согласования ФСБ, др. спец.права",
    "Окончание срока действия трудового договора",
    "Выход на пенсию",
    TECHNICAL_REASON,
}

# Статья ТК РФ по причине; для остальных — «по собственному желанию»
REASON_ARTICLE = {
    "Окончание срока действия трудового договора": "пункт 2 части первой статьи 77",
    TECHNICAL_REASON: "пункт 2 части первой статьи 77",
}
DEFAULT_ARTICLE = "пункт 3 части первой статьи 77"


def write_employees_report(
    rows: list[FakeRow], path: Path, *, mislabel_fire_date: bool = False
) -> None:
    """Файл «Личные данные сотрудников»: шапка, заголовок в 9-й строке, данные с 10-й.

    mislabel_fire_date=True повторяет ловушку настоящей выгрузки: колонка с датами
    увольнения называется «Адрес места проживания».
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Лист_1"

    people = len({row.episode.snils for row in rows})  # сколько разных СНИЛС
    ws["A2"] = "Личные данные сотрудников"
    ws["A4"] = "Отбор:"
    ws["B4"] = f'Организация Равно "{ORGANIZATION}"'
    ws["A6"] = "Организация"
    ws["E6"] = ORGANIZATION
    ws["A7"] = "Количество"
    ws["E7"] = format_count(people)

    header = list(EMPLOYEES_HEADER)
    if mislabel_fire_date:
        header[header.index("Дата увольнения")] = MISLABELED_FIRE_DATE
    for col, title in enumerate(header, start=1):
        ws.cell(row=EMPLOYEES_HEADER_ROW, column=col, value=title)

    for offset, row in enumerate(rows, start=1):
        e = row.episode
        values = (
            e.tab_number,
            None,
            None,
            row.full_name,
            None,
            e.position,
            ru_date(e.hire_date),
            ru_date(e.fire_date) or None,  # пустая дата → пустая ячейка
            row.staff_category,
            e.snils,
            row.unit.department,
            row.unit.top_department,
            row.unit.functional_subordination,
            row.unit.functional_block,
        )
        for col, value in enumerate(values, start=1):
            ws.cell(row=EMPLOYEES_HEADER_ROW + offset, column=col, value=value)

    wb.save(path)


def write_reasons_report(rows: list[FakeRow], path: Path, *, period: str) -> None:
    """Файл «Причины увольнений»: двухуровневый заголовок в 4–5-й строках, данные с 6-й."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Лист_1"
    ws["A2"] = "Параметры:"
    ws["C2"] = f"Период: {period}"

    top = (
        "Организация",
        "Дата регистрации",
        "Должность",
        "Причина по мнению HR",
        "Причина по мнению руководителя",
        "Увольнение.Дата увольнения",
        "Прием на работу.Дата приема",
        "Увольнение.Статья ТК РФ",
        "сотрудник",  # над двумя колонками
        None,
        "группа",  # над четырьмя колонками
        None,
        None,
        None,
    )
    sub = (None,) * 8 + (
        "Сотрудник.Таб. номер",
        "Сотрудник",
        "Подразделение",
        "Подразделение.Функциональное подчинение (Подразделения)",
        "Подразделение.Функциональный блок (Подразделения)",
        "Подразделение.Подразделение верхнего уровня (для отчетов) (Подразделения)",
    )
    for col, (title, subtitle) in enumerate(zip(top, sub, strict=True), start=1):
        ws.cell(row=4, column=col, value=title)
        ws.cell(row=5, column=col, value=subtitle)
    ws.merge_cells("I4:J4")  # «сотрудник» над двумя колонками
    ws.merge_cells("K4:N4")  # «группа» над четырьмя

    fired = [row for row in rows if row.episode.fire_reason is not None]
    fired.sort(key=lambda row: row.episode.fire_date)
    for r, row in enumerate(fired, start=6):
        e = row.episode
        values = (
            ORGANIZATION,
            None,
            e.position,
            e.fire_reason,
            e.fire_reason,
            ru_date(e.fire_date),
            ru_date(e.hire_date),
            REASON_ARTICLE.get(e.fire_reason, DEFAULT_ARTICLE),
            e.tab_number,
            row.full_name,
            row.unit.department,
            row.unit.functional_subordination,
            row.unit.functional_block,
            row.unit.top_department,
        )
        for col, value in enumerate(values, start=1):
            ws.cell(row=r, column=col, value=value)

    wb.save(path)


def write_reasons_reference(path: Path) -> None:
    """Справочник причин: формулировка ЗАГЛАВНЫМИ и признак «Входит / Не входит»."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Причины"
    ws.append(["Наименование", "Учёт в текучести"])
    all_reasons = [reason for reason, _ in REASONS] + [TECHNICAL_REASON]
    for reason in sorted(all_reasons):
        status = "Не входит" if reason in EXCLUDED_REASONS else "Входит"
        ws.append([reason.upper(), status])
    wb.save(path)
