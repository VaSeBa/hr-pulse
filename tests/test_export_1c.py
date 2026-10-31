"""Тесты для hr_pulse.demo.export_1c: файлы похожи на выгрузки 1С и согласованы между собой."""

import re

import pytest
from openpyxl import load_workbook

from hr_pulse.demo.export_1c import (
    EMPLOYEES_HEADER_ROW,
    MISLABELED_FIRE_DATE,
    write_employees_report,
    write_reasons_reference,
    write_reasons_report,
)
from hr_pulse.demo.generator import FakeCompany, GeneratorConfig

DATE_RE = re.compile(r"^\d{2}\.\d{2}\.\d{4}$")  # дд.мм.гггг


@pytest.fixture(scope="module")
def rows():
    return FakeCompany(GeneratorConfig(people=300, seed=5)).generate()


@pytest.fixture(scope="module")
def employees_sheet(rows, tmp_path_factory):
    path = tmp_path_factory.mktemp("export") / "employees.xlsx"
    write_employees_report(rows, path)
    return load_workbook(path).active


@pytest.fixture(scope="module")
def reasons_sheet(rows, tmp_path_factory):
    path = tmp_path_factory.mktemp("export") / "reasons.xlsx"
    write_reasons_report(rows, path, period="01.01.2024 - 30.09.2026")
    return load_workbook(path).active


def header_of(sheet, row_number):
    return [cell.value for cell in sheet[row_number]]


def test_employees_title_and_count(employees_sheet, rows):
    assert employees_sheet["A2"].value == "Личные данные сотрудников"
    people = len({r.episode.snils for r in rows})
    assert employees_sheet["E7"].value.replace("\xa0", "") == str(people)


def test_employees_header_in_row_9(employees_sheet):
    header = header_of(employees_sheet, EMPLOYEES_HEADER_ROW)
    assert header[0] == "Табельный номер (с префиксами)"
    assert header[1] is None  # пустая служебная колонка
    assert "СНИЛС" in header
    assert "Дата увольнения" in header


def test_employees_one_line_per_episode(employees_sheet, rows):
    assert employees_sheet.max_row == EMPLOYEES_HEADER_ROW + len(rows)


def test_employees_dates_are_text(employees_sheet):
    header = header_of(employees_sheet, EMPLOYEES_HEADER_ROW)
    hire_col = header.index("Дата приема") + 1
    values = [
        employees_sheet.cell(row=r, column=hire_col).value
        for r in range(EMPLOYEES_HEADER_ROW + 1, employees_sheet.max_row + 1)
    ]
    assert all(isinstance(v, str) and DATE_RE.match(v) for v in values)


def test_mislabeled_fire_date(rows, tmp_path):
    path = tmp_path / "trap.xlsx"
    write_employees_report(rows, path, mislabel_fire_date=True)
    header = header_of(load_workbook(path).active, EMPLOYEES_HEADER_ROW)
    assert MISLABELED_FIRE_DATE in header
    assert "Дата увольнения" not in header


def test_reasons_two_level_header(reasons_sheet):
    assert reasons_sheet["I4"].value == "сотрудник"
    assert reasons_sheet["I5"].value == "Сотрудник.Таб. номер"
    merged = {str(rng) for rng in reasons_sheet.merged_cells.ranges}
    assert {"I4:J4", "K4:N4"} <= merged


def test_reasons_only_dismissals_with_reason(reasons_sheet, rows):
    with_reason = [r for r in rows if r.episode.fire_reason is not None]
    assert reasons_sheet.max_row - 5 == len(with_reason)


def test_reasons_join_to_employees_by_tab_number(reasons_sheet, rows):
    # Каждая строка причин находит свой эпизод по табельному номеру — как в настоящих файлах
    tabs = {r.episode.tab_number for r in rows}
    reason_tabs = [
        reasons_sheet.cell(row=r, column=9).value for r in range(6, reasons_sheet.max_row + 1)
    ]
    assert reason_tabs
    assert all(tab in tabs for tab in reason_tabs)


def test_reference_statuses(tmp_path):
    path = tmp_path / "reference.xlsx"
    write_reasons_reference(path)
    sheet = load_workbook(path).active
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    statuses = {status for _, status in rows}
    assert statuses == {"Входит", "Не входит"}
    assert all(name == name.upper() for name, _ in rows)  # заглавными, как в справочнике компании
